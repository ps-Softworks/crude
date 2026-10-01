"""Tests für das Tabellenmodell (Roadmap 0.6).

Baut die Excel-Datei frisch, setzt je Szenario Eingaben, lässt alle Formeln
mit der Bibliothek `formulas` rechnen und vergleicht mit einer unabhängigen
Python-Nachrechnung der Spielregeln.

Aufruf:  .venv/bin/python -m unittest tools/tabellenmodell/test_tabellenmodell.py
"""

import sys
import tempfile
import unittest
from pathlib import Path

import formulas
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_tabellenmodell as bt  # noqa: E402

S = bt.SP
STANDARD = {name: wert for name, _t, wert, *_r in bt.STARTWERTE}
STANDARD.update({name: wert for name, _t, wert, *_r in bt.KREDIT})
LAGEN = {lage: (bonus, roy) for lage, bonus, roy, _q in bt.LAGEN}
BESITZER = {b: (faktor, zuschlag) for b, faktor, zuschlag, _w in bt.BESITZER}
RATING = dict(bt.RATINGS)
ROY_MIN, ROY_MAX = 0.10, 0.25


def runde(pachten=1, bohren=1, ergebnis="Trocken", groesse=0, lage="Randlage",
          besitzer="neutral", bonus_v=None, roy_v=None):
    return dict(pachten=pachten, bohren=bohren, ergebnis=ergebnis, groesse=groesse, lage=lage,
                besitzer=besitzer, bonus_v=bonus_v, roy_v=roy_v)


PAUSE = runde(pachten=0, bohren=0)
PECHSTRAEHNE = [runde()] * bt.RUNDEN


# --- Unabhängige Nachrechnung -------------------------------------------
def verhandle(e):
    bonus, roy = LAGEN[e["lage"]]
    faktor, zuschlag = BESITZER[e["besitzer"]]
    b = e["bonus_v"] if e["bonus_v"] is not None else bonus * faktor
    f = e["roy_v"] if e["roy_v"] is not None else min(ROY_MAX, max(ROY_MIN, roy + zuschlag))
    return b, f


def bankzins(p, quellen):
    if p["bank_verhandelt"] is not None:
        return p["bank_verhandelt"]
    z = RATING[p["rating"]] + (-p["sicher_rabatt"] if quellen > 0 else p["ohne_aufschlag"])
    return max(0, z + p["klima"])


def referenz(p, runden):
    bar, bank, leiher = p["start_cash"], 0.0, 0.0
    gepachtet, trocken, pleite, fund_nr = 0, 0, False, 0
    quellen = []
    zeilen = []
    for r, e in enumerate(runden, start=1):
        bonus, roy = verhandle(e)
        alt_bezahlt = max(0, gepachtet - p["free_leases"])
        gepachtet += e["pachten"]
        k_pacht = (max(0, gepachtet - p["free_leases"]) - alt_bezahlt) * bonus
        k_bohr = e["bohren"] * p["drill_cost"]

        mengen = []
        for q in quellen:
            alter = r - q["runde"] - 1
            menge = 0.0
            if alter >= 0 and q["r0"] > 0:
                menge = min(q["rest"], p["q0"] * p["days"] * (1 - p["decline"]) ** alter
                            * q["rest"] / q["r0"])
            mengen.append(menge)
        fördernd = sum(1 for m in mengen if m > 0)
        bz = bankzins(p, fördernd)
        rahmen = p["bank_grund"] + fördernd * p["bank_je_quelle"]
        zinsen = bank * bz / 4 + leiher * p["leiher_zins"] / 4

        foerd = foerderzins = silas = 0.0
        for q, m in zip(quellen, mengen):
            q["rest"] -= m
            foerd += m
            foerderzins += m * p["price"] * q["roy"]
            if q["silas"]:
                silas += m * (p["price"] - p["transport"] - p["price"] * q["roy"]) \
                    * p["silas_share"]
        erloes = foerd * (p["price"] - p["transport"])
        vor = bar + erloes - foerderzins - silas - k_pacht - k_bohr - zinsen

        # Kredit: billigeres Geld zuerst, teureres zuerst tilgen
        bedarf = max(0, -vor)
        bank_frei = max(0, rahmen - bank)
        leiher_frei = max(0, p["leiher_rahmen"] * p["leiher_an"] - leiher)
        geber = [["bank", bz, bank_frei], ["leiher", p["leiher_zins"], leiher_frei]]
        geber.sort(key=lambda g: (g[1], g[0] != "bank"))  # bei Gleichstand Bank zuerst
        neu = {"bank": 0.0, "leiher": 0.0}
        for name, _z, frei in geber:
            neu[name] = min(bedarf, frei)
            bedarf -= neu[name]
        ueber = max(0, vor) if p["auto_repay"] == 1 else 0
        schuld = {"bank": bank, "leiher": leiher}
        tilg = {"bank": 0.0, "leiher": 0.0}
        for name, _z, _f in reversed(geber):
            tilg[name] = min(ueber, schuld[name])
            ueber -= tilg[name]
        bank += neu["bank"] - tilg["bank"]
        leiher += neu["leiher"] - tilg["leiher"]
        bar = vor + neu["bank"] + neu["leiher"] - tilg["bank"] - tilg["leiher"]

        if e["bohren"] == 1 and e["ergebnis"] == "Trocken":
            trocken += 1
        pleite = pleite or bar < 0
        if e["bohren"] == 1 and e["ergebnis"] == "Fund":
            fund_nr += 1
            quellen.append(dict(runde=r, r0=e["groesse"], rest=e["groesse"], roy=roy,
                                silas=fund_nr <= p["silas_wells"]))
        zeilen.append(dict(bonus=bonus, roy=roy, bankzins=bz, bankrahmen=rahmen,
                           zinsen=zinsen, foerd=foerd, foerderzins=foerderzins, silas=silas,
                           bank_neu=neu["bank"], leiher_neu=neu["leiher"],
                           bank_tilg=tilg["bank"], leiher_tilg=tilg["leiher"],
                           bank_end=bank, leiher_end=leiher, bar_end=bar, trocken=trocken,
                           status="PLEITE" if pleite else "ok"))
    return zeilen


def referenz_pleiterechner(p, kosten, rahmen):
    zins = bankzins(p, 0)
    bonus = LAGEN["Randlage"][0] * BESITZER["neutral"][0]
    summe, bezahlbar = 0.0, 0
    for r in range(1, bt.RUNDEN + 1):
        z = max(0, summe - p["start_cash"]) * zins / 4 if r > 1 else 0
        summe += kosten + (bonus if r > p["free_leases"] else 0) + z
        if summe <= p["start_cash"] + rahmen:
            bezahlbar += 1
    return bezahlbar


# --- Tabelle rechnen lassen ----------------------------------------------
def rechne(werte=None, runden=None):
    tmp = Path(tempfile.mkdtemp()) / "modell.xlsx"
    bt.baue(tmp)
    wb = load_workbook(tmp)
    werte = werte or {}
    for i, (name, *_r) in enumerate(bt.STARTWERTE):
        if name in werte:
            wb["Startwerte"][f"B{bt.ERSTE_STARTZEILE + i}"] = werte[name]
    for i, (name, *_r) in enumerate(bt.KREDIT):
        if name in werte:
            wb["Kreditgeber"][f"B{bt.KG0 + i}"] = werte[name]
    if runden:
        ws = wb["Runden"]
        for r, e in enumerate(runden, start=1):
            z = bt.rz(r)
            for key, sp in [("pachten", "pacht"), ("bohren", "bohren"), ("ergebnis", "ergebnis"),
                            ("groesse", "groesse"), ("lage", "lage"), ("besitzer", "besitzer"),
                            ("bonus_v", "bonus_v"), ("roy_v", "roy_v")]:
                ws[f"{S[sp]}{z}"] = e[key]
    wb.save(tmp)
    sol = formulas.ExcelModel().loads(str(tmp)).finish().calculate()
    ergebnis = {}
    for k, v in sol.items():
        if "!" in k and ":" not in k:
            blatt, zelle = k.split("]", 1)[1].split("!")
            ergebnis[(blatt.strip("'").upper(), zelle)] = v.value[0, 0]

    def get(blatt, zelle):
        return ergebnis[(blatt.upper(), zelle)]
    return get


class Basis(unittest.TestCase):
    P = STANDARD
    RUNDEN = PECHSTRAEHNE
    WERTE = None

    @classmethod
    def setUpClass(cls):
        cls.get = staticmethod(rechne(cls.WERTE, cls.RUNDEN))
        cls.ref = referenz(cls.P, cls.RUNDEN)

    def wert(self, key, r):
        return self.get("Runden", f"{S[key]}{bt.rz(r)}")

    def ergebnis(self, i):
        return self.get("Runden", f"G{bt.ergebnis_zeile() + i}")

    def test_alle_runden_wie_referenz(self):
        for r, erw in enumerate(self.ref, start=1):
            for key, soll in erw.items():
                with self.subTest(runde=r, spalte=key):
                    if isinstance(soll, str):
                        self.assertEqual(self.wert(key, r), soll)
                    else:
                        self.assertAlmostEqual(self.wert(key, r), soll, places=4)


class TestStandard(Basis):
    """Vorgabe: Pechsträhne in Randlage mit den Startwerten aus GDD §15."""

    def test_startwerte_aus_gdd(self):
        self.assertEqual(STANDARD["start_cash"], 2000)
        self.assertEqual(LAGEN["Randlage"][1], 0.125)
        self.assertEqual(STANDARD["silas_share"], 0.30)
        self.assertEqual(RATING, {"A": 0.05, "B": 0.07, "C": 0.10, "D": 0.15})
        self.assertTrue(1500 <= STANDARD["drill_cost"] <= 3000)
        self.assertTrue(50 <= LAGEN["Randlage"][0] <= 300)
        self.assertTrue(all(1000 <= LAGEN[l][0] <= 20000
                            for l in ("Nachbar eines Funds", "Am Fund")))
        self.assertTrue(0.08 <= STANDARD["decline"] <= 0.15)
        self.assertTrue(0.10 <= STANDARD["hit_rate"] <= 0.20)

    def test_pleite_nach_drei_trockenbohrungen(self):
        # R1: 2.000 $ bar. R2: 2.000 $ Bank. R3: 2.150 $ + 65 $ Zins (13 % auf 2.000 $):
        # 1.000 $ Bank-Rest + 1.215 $ Geldverleiher. R4: 785 $ Verleiher-Rest reicht nicht.
        self.assertEqual(self.ergebnis(0), 4)
        self.assertEqual(self.ergebnis(1), 3)
        self.assertAlmostEqual(self.wert("leiher_neu", 3), 1215)
        self.assertAlmostEqual(self.wert("zinsen", 4), 3000 * 0.13 / 4 + 1215 * 0.40 / 4)

    def test_bank_ohne_sicherheit_teurer(self):
        self.assertAlmostEqual(self.wert("bankzins", 1), 0.13)

    def test_pacht_erst_nach_freien_optionen(self):
        self.assertEqual(self.wert("k_pacht", 2), 0)
        self.assertEqual(self.wert("k_pacht", 3), 150)

    def test_pleiterechner_wie_referenz(self):
        for i, kosten in enumerate(bt.PR_KOSTEN):
            for k, rahmen in enumerate(bt.PR_RAHMEN):
                n = referenz_pleiterechner(STANDARD, kosten, rahmen)
                with self.subTest(kosten=kosten, rahmen=rahmen):
                    self.assertEqual(self.get("Pleiterechner", f"{'CDEF'[k]}{6 + i}"), n)
                    self.assertAlmostEqual(self.get("Pleiterechner", f"{'CDEF'[k]}{13 + i}"),
                                           (1 - STANDARD["hit_rate"]) ** n)

    def test_pleiterechner_ohne_kredit_eine_bohrung(self):
        self.assertEqual(self.get("Pleiterechner", "C7"), 1)


class TestPachtverhandlung(Basis):
    """Lage und Landbesitzer bestimmen Bonus und Förderzins; Verhandeltes überschreibt."""

    RUNDEN = [
        runde(lage="Am Fund", besitzer="gierig"),
        runde(lage="Randlage", besitzer="verschuldet"),
        runde(lage="Nachbar eines Funds", besitzer="misstrauisch"),
        runde(lage="Nachbar eines Funds", besitzer="fromm", bonus_v=500, roy_v=0.15),
    ] + [PAUSE] * (bt.RUNDEN - 4)
    WERTE = dict(bank_grund=50000)

    P = dict(STANDARD, bank_grund=50000)

    def test_gieriger_besitzer_am_fund(self):
        self.assertAlmostEqual(self.wert("bonus", 1), 8000 * 1.25)
        self.assertAlmostEqual(self.wert("roy", 1), 0.23)

    def test_verschuldeter_besitzer_untergrenze(self):
        # 1/8 − 3 Punkte = 9,5 % liegt unter der Untergrenze 10 %
        self.assertAlmostEqual(self.wert("roy", 2), 0.10)

    def test_misstrauischer_besitzer_will_bonus(self):
        self.assertAlmostEqual(self.wert("bonus", 3), 2000 * 1.5)
        self.assertAlmostEqual(self.wert("k_pacht", 3), 3000)  # 3. Parzelle, Optionen weg

    def test_verhandeltes_ueberschreibt(self):
        self.assertEqual(self.wert("bonus", 4), 500)
        self.assertEqual(self.wert("roy", 4), 0.15)


class TestFundMitSicherheit(Basis):
    """Fund am Fund-Rand mit gierigem Besitzer, dann Bohrungen auf Kredit mit Pfand."""

    RUNDEN = ([runde(pachten=0, ergebnis="Fund", groesse=200000, lage="Nachbar eines Funds",
                     besitzer="gierig")]
              + [runde(ergebnis="Fund", groesse=20000, besitzer="verschuldet")]
              + [runde()] * 3 + [PAUSE] * (bt.RUNDEN - 5))

    def test_foerderzins_je_quelle(self):
        # Runde 3: Quelle 1 (1/6 + 3 Punkte), Quelle 2 (Untergrenze 10 %)
        q1 = self.get("Quellen", f"B{bt.q_zeile(3)}")
        q2 = self.get("Quellen", f"C{bt.q_zeile(3)}")
        self.assertGreater(q2, 0)
        self.assertAlmostEqual(self.wert("foerderzins", 3), q1 * (1 / 6 + 0.03) + q2 * 0.10)

    def test_pfand_macht_bank_billiger_und_rahmen_groesser(self):
        self.assertAlmostEqual(self.wert("bankzins", 1), 0.13)
        self.assertAlmostEqual(self.wert("bankzins", 2), 0.08)
        self.assertEqual(self.wert("quellen", 3), 2)
        self.assertEqual(self.wert("bankrahmen", 3), 3000 + 2 * 2000)


class TestGeldverleiherAus(Basis):
    P = dict(STANDARD, leiher_an=0)
    WERTE = dict(leiher_an=0)

    def test_pleite_nach_zwei_trockenbohrungen(self):
        self.assertEqual(self.ergebnis(0), 3)
        self.assertEqual(self.ergebnis(1), 2)
        self.assertEqual(self.wert("leiher_neu", 3), 0)


class TestBankTeurerAlsVerleiher(Basis):
    """Schlecht verhandelt: 50 % Bankzins – dann nimmt Jacob zuerst den Geldverleiher."""

    P = dict(STANDARD, bank_verhandelt=0.50)
    WERTE = dict(bank_verhandelt=0.50)

    def test_verleiher_zuerst(self):
        self.assertAlmostEqual(self.wert("bankzins", 1), 0.50)
        self.assertEqual(self.wert("leiher_neu", 2), 2000)
        self.assertEqual(self.wert("bank_neu", 2), 0)


class TestTilgung(Basis):
    """Fund, Kredite bei beiden Geldgebern, automatische Tilgung: teureres zuerst."""

    RUNDEN = ([runde(), runde(), runde(ergebnis="Fund", groesse=2000000), runde()]
              + [PAUSE] * (bt.RUNDEN - 4))
    P = dict(STANDARD, auto_repay=1, q0=400)
    WERTE = dict(auto_repay=1, q0=400)

    def test_teureres_zuerst_getilgt(self):
        tilg = [(r, self.wert("leiher_tilg", r), self.wert("bank_tilg", r))
                for r in range(1, bt.RUNDEN + 1)]
        erste_leiher = next(r for r, l, _b in tilg if l > 0)
        erste_bank = next(r for r, _l, b in tilg if b > 0)
        self.assertLessEqual(erste_leiher, erste_bank)
        self.assertEqual(self.wert("leiher_end", bt.RUNDEN), 0)
        self.assertEqual(self.wert("bank_end", bt.RUNDEN), 0)


if __name__ == "__main__":
    unittest.main()
